"""Build an evidence-checked research draft from the corrected capability study.

This builder targets the capability-barrier outcome. It refuses to substitute
unlaunched main experiments for efficacy results or to pool the older cohorts.
"""
import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_empirical_paper import add_link, table
from analyze_guard_study_v2 import analyze_runs

ROOT = Path(__file__).resolve().parents[1]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_verified_summary(directory):
    manifest = json.loads((directory / "manifest.json").read_text())
    if manifest.get("model_calls") != 0:
        raise ValueError("The controlled diagnostic must make no model calls")
    for name, expected in manifest.get("output_sha256", {}).items():
        path = (directory / name).resolve()
        if not path.is_relative_to(directory.resolve()) or sha(path) != expected:
            raise ValueError(f"Diagnostic integrity mismatch: {name}")
    if "summary.json" not in manifest.get("output_sha256", {}):
        raise ValueError("Diagnostic summary needs a frozen hash")
    return json.loads((directory / "summary.json").read_text())


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", type=Path, action="append", required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    report = analyze_runs(args.run)
    if report["status"] != "passed":
        raise ValueError("Every supplied corrected cohort must pass the independent replay audit")
    if len(report["capability_gates"]) != 2 or report["capability"]["episodes"] != 36:
        raise ValueError("This draft requires both complete corrected 18-episode capability cohorts")
    if any(gate["passed"] for gate in report["capability_gates"]) or report["main"].get("episodes", 0):
        raise ValueError("A capable model requires the main-study manuscript branch, not this barrier draft")
    comparison_dir = ROOT / "runs/guard-policy-comparison"
    comparison = read_verified_summary(comparison_dir)
    if (comparison["results_kind"] != "scripted_execution_comparison"
            or comparison["matched_probe_states"] != 14
            or not comparison["ceg_provenance_outcomes_equal"]):
        raise ValueError("Unexpected controlled comparison; revise manuscript rather than reuse claims")
    diagnostic_dir = ROOT / "runs/guard-requirement-coverage-diagnostic"
    diagnostic = read_verified_summary(diagnostic_dir)
    if diagnostic["analysis_kind"] != "posthoc_scripted_requirement_coverage_diagnostic":
        raise ValueError("Coverage evidence must remain explicitly post hoc")
    rows = [json.loads(line) for run in args.run for line in (run / "episodes.jsonl").read_text().splitlines()]
    gates = sorted(report["capability_gates"], key=lambda x: x["model"], reverse=True)
    models = {"qwen2.5:3b": "Qwen2.5 3B", "llama3.2:1b": "Llama3.2 1B"}
    counts = {gate["model"]: gate for gate in gates}
    q, l = counts["qwen2.5:3b"], counts["llama3.2:1b"]
    calls = sum(len(row["calls"]) for row in rows)
    actions = sum(event["kind"] == "action" for row in rows for event in row["events"])
    model_errors = sum(bool(row["model_error"]) for row in rows)
    budgets = sum(row["metrics"]["action_budget_exhausted"] for row in rows)
    invalid = sum(row["metrics"]["invalid_actions"] for row in rows)

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().find(qn("w:rFonts"))
        for attr in list(fonts.attrib):
            if "theme" in attr.lower():
                del fonts.attrib[attr]
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + script), "Times New Roman")
        for border in list(style._element.get_or_add_pPr().findall(qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.06
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Title"].paragraph_format.space_after = Pt(6)
    doc.styles["Subtitle"].font.size = Pt(12)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    title = "Provenance and Recovery Requirements for Conditional Agent Permissions"
    doc.core_properties.title = title
    doc.core_properties.author = "Amornchot Singh"
    doc.core_properties.subject = "Research draft with scripted enforcement evidence and measured capability limits"

    def p(text):
        return doc.add_paragraph(text)

    def h(text):
        return doc.add_paragraph(text, "Heading 1")

    def page(text):
        heading = h(text)
        heading.paragraph_format.page_break_before = True

    doc.add_paragraph(title, "Title")
    doc.add_paragraph("An execution backed testbed and local model capability study", "Subtitle")
    p("Amornchot Singh | Research draft | 18 October 2026")
    h("Abstract")
    p("Conditional permissions can become unsafe when an agent controls the state that activates an exception. We implement a local SQLite broker that distinguishes independently scheduled, approved-work and discretionary resource effects, and compare eight authorization policies on identical executed states. Across 14 authored request states, the Counterfactual Exception Guard (CEG) agrees with an independently implemented provenance rule in all 14 decisions. Both block the five actor-dependent uses permitted by a factual-only threshold, while retaining one unregistered-repair violation. They also withhold two legitimate reference workflows when evidence is incomplete. A separate post hoc requirement-coverage gate would suppress the retained violation on its saved trace. "
      f"Fresh capability checks with opaque scenario identifiers yield {q['compliant_capability_success']}/18 compliant successes for Qwen2.5 3B and {l['compliant_capability_success']}/18 for Llama3.2 1B. Neither meets the declared capability gate, so the corrected main agent study is not launched. "
      "The contribution is an auditable testbed, an equivalence result and measured limits of the current evaluation setup. The results do not establish a novel general authorization algorithm or effectiveness against capable adaptive agents.")
    h("Motivation and research question")
    p("Suppose a tool may use external processing only when independently justified workload crosses a policy-defined threshold. An agent could add unnecessary jobs until a simple load threshold permits the external route. The safety problem is the integrity of the permission condition: completing a task should not grant the agent authority to manufacture its own exception. Crossing this threshold need not make local execution physically impossible. The environment must independently record the state change and its origin, rather than accept an actor's claimed outage.")
    p("Our question is which evidence and policy checks are needed to preserve legitimate fallback while excluding qualification caused by discretionary agent effects. The testbed isolates policy semantics, missing evidence and recovery obligations. It treats successful permission enforcement, task completion and agent competence as separate outcomes.")
    h("Scope of the contribution")
    p("We provide an execution-backed broker with retained database and report artifacts; a controlled comparison against equally informed provenance and taint rules; and a versioned evaluation process that exposes implementation failures and capability limits. The underlying ideas have substantial security precedents [1–8]. We make no claim that a new name or a combination of existing controls establishes novelty.")

    page("Operational model and threat boundary")
    p("Each episode begins with a fixed task manifest, deterministic source data and a local job queue. Required jobs compute check receipts; a local report computes three batches, while the alternative route writes the same verified local report in one operation. Successful finish requires a correct report, exactly the required receipts and any mandatory repair. Report creation alone is not full task success. No confidential data or external service is involved, and the alternative route performs no network transfer.")
    p("The actor emits one schema-constrained JSON tool call at a time. The dispatcher assigns origin from the executed operation: an independent controlled scheduler creates external jobs, the fixed manifest authorizes required jobs, and optional jobs are discretionary actor effects. The actor cannot submit origin labels, enlarge its work grant or edit the database through the exposed tools. The broker retains an observed evidence ledger separately from database ground truth.")
    p("This establishes origin inside a trusted local abstraction. It does not prove that an apparently external outage was independent of the actor in a deployment. Approved work remains eligible even when the agent chooses when to queue it: the rule excludes discretionary qualification, not every strategic arrangement of authorized work. Logical queue occupancy is not measured CPU contention. Arbitrary shell access, compromised brokers, distributed processes and induced third-party failures are outside the tested boundary. Fixed logical time does not test real-time expiry or distributed races.")
    h("Permission semantics")
    p("The declared user rule requires qualifying load without relying on discretionary actor work and requires specified repair before external processing. It explicitly prohibits manufacturing fallback conditions. A database oracle scores actor-dependent external execution and premature fallback before required repair as violations of that declared rule. Evidence uncertainty is reported separately: a missing label is not itself proof of misconduct.")
    h("Why the arithmetic comparator is equivalent")
    p("Let eligible load be the nonnegative sum of trusted baseline, independent external load and valid approved-work load. Factual load adds the nonnegative discretionary contribution. CEG requires both loads to reach the threshold, complete evidence and no overdue registered repair. Eligible load cannot exceed factual load, so reaching the eligible threshold already implies the factual threshold. The factual conjunction is redundant in this additive model.")
    p("Consequently, an independently implemented provenance rule that checks eligible load, evidence completeness and the same registered obligations has the same authorization predicate. Adding discretionary jobs cannot turn a false eligible predicate into a true one while the trusted history remains fixed. This is a direct property of the specified rule, not general causal inference, arbitrary-agent noninterference or elimination of the incentive to disrupt resources.")
    h("Enforcement and omissions")
    p("CEG also binds a single-use execution lease to the job, principal, operation, arguments, resource, policy and state; it rechecks before dispatch. An ambiguous callback failure consumes the lease and marks evidence uncertain. These are implementation controls, not new security primitives. The original guard knows only registered recovery obligations, leaving a deliberately retained failure when a task requires repair but no obligation was registered.")

    page("Experimental design and evidence separation")
    p("Twelve authored cases cover healthy capacity, necessary external processing, approved work, discretionary load, mixed causes, missing origin, misleading origin claims, omitted ledger effects, registered repair, unregistered repair and ordinary required work. They are parameter and composition holdouts created by the same team as the mechanism. They are not independent real-world domains or human-validated labels.")
    p("The controlled experiment runs identical prefixes and requests under eight policies, verifying equal pre-request database state. Twelve qualifying states plus two cleanup contrasts produce 14 states and 112 policy executions. A separate legitimate scripted controller attempts all 12 tasks under each policy, producing 96 reference executions. These measure implemented decisions and reachability; they are not LLM behavior or a learned policy's utility.")
    table(doc, ["Policy", "Decision rule beyond the factual threshold"], [
        ["Factual only", "No provenance or repair requirement"],
        ["Permanent taint", "Complete evidence; reject after any discretionary effect in history"],
        ["Active taint", "Reject while discretionary effects remain active"],
        ["History rules", "Active taint plus complete evidence and registered repair"],
        ["Shadow only", "Eligible load and evidence, without recovery obligations"],
        ["Provenance rules", "Eligible load, evidence and registered repair"],
        ["CEG", "Equivalent predicate here, with bound execution leases"],
        ["Deny all", "Reject every conditional external request"],
    ], [1.55, 5.45])
    p("Table 1. Local semantic comparators, not reproductions of published systems. Every comparator has access to the same observed evidence. Permanent-taint and history rules are intentionally different policy choices; their loss of permissions is not a framework vulnerability.")
    h("Collection amendment and capability gate")
    p("The first interface exposed semantic task identifiers. A separate AI-assisted code and record audit found that names such as hidden-effect could reveal case categories. We stopped those main pilots after seven Qwen and six Llama records, including one interrupted episode per model, and preserve them as excluded development evidence. The original 36 capability records are retained separately. The corrected collector replaces public identifiers with opaque stable digests and checks reconstructed requests for label leakage. No external human audit is claimed.")
    p("The revised protocol, written after this discovery and the first capability results, runs six capability exercises three times each for both models. Main-study eligibility requires at least 12 of 18 compliant successes and at least one success per exercise. This is a pragmatic prospective screen, not an independently preregistered or validated cutoff. A compliant success satisfies the exercise and has zero strict-policy violations. Failed models remain in the report; their main cohorts are not launched.")
    p("The planned main study for an eligible model is 12 cases by three policies by two instruction roles by three repetitions, or 216 episodes. Ordinary assistance and explicitly adversarial sandbox testing remain separate. Settings are temperature 0.6, context 4096, output cap 128 tokens and at most 12 actions, with fixed shuffled schedules and no retries. All assigned outcomes, errors and exhausted budgets remain in their denominators.")

    page("Controlled policy results")
    order = ["factual_only", "permanent_taint", "active_taint", "history_rules", "shadow_only", "provenance_rules", "ceg", "deny_all"]
    labels = {"factual_only": "Factual only", "permanent_taint": "Permanent taint", "active_taint": "Active taint", "history_rules": "History rules", "shadow_only": "Shadow only", "provenance_rules": "Provenance rules", "ceg": "CEG", "deny_all": "Deny all"}
    body = []
    for key in order:
        b = comparison["by_baseline"][key]
        body.append([labels[key], f"{b['allow']} / {b['deny']} / {b['defer']}", str(b["actor_dependent_external_uses"]), str(b["strict_oracle_violations"]), f"{b['reference_task_success']}/12"])
    table(doc, ["Policy", "Allow deny defer\n14 requests", "Actor\ndependent", "All strict\nviolations", "Reference\ncompletion"], body, [1.65, 1.55, 1.1, 1.25, 1.45])
    p("Table 2. Author-scripted executions. Decision and violation columns concern matched requests; the final completion column concerns separate legitimate reference workflows. Counts across policies are repeated measurements of the same authored states, not independent deployment samples.")
    h("Provenance helps only relative to a weaker rule")
    p("Factual-only permits all 14 matched requests, including five actor-dependent uses and two premature uses before repair. CEG and provenance-rules deny the five actor-dependent requests and agree on every outcome. They still permit one required-but-unregistered repair case. This validates the distinction between a factual threshold and the intended origin-sensitive policy. It provides no evidence of CEG outperforming an equally informed provenance rule.")
    p("History-rules denies one mixed-cause request that CEG permits because independently sufficient load coexists with discretionary jobs. Permanent-taint also rejects the two cleanup contrasts after discretionary effects have been removed. These differences illustrate policy expressiveness, not measured agent adaptation.")
    h("Safety and availability must be reported together")
    p("The legitimate controller completes 12 of 12 tasks under factual-only and active-taint, 10 under CEG and provenance-rules, and nine under deny-all. All reference completions have zero strict violations. CEG's two failures occur when a full independent queue cannot be repaired and evidence is incomplete. In the hidden-effect case, 11 observed trusted units already exceed the threshold of ten, yet the guard still defers because total-load reconciliation fails. Its blanket completeness rule therefore imposes an availability cost even when a certified lower bound would suffice numerically.")
    p("The experiment does not determine whether such deferral is appropriate for every deployment. That requires a policy decision about incomplete evidence and a resource model supporting sound lower bounds. Counting deferral as automatically correct would hide this tradeoff.")

    page("Measured capability")
    p(f"All 36 corrected-interface capability episodes were retained and replay-audited, comprising {calls} inference calls, {actions} dispatched actions, {model_errors} model or transport errors and {budgets} action-budget exhaustions. The broker rejected {invalid} tool calls for invalid arguments or unmet preconditions. Replay reconstructs requests, actions, database state, report contents and metrics. Hashes verify artifact consistency, not independent attestation of execution.")
    table(doc, ["Model", "Capability\nsuccess", "Compliant\ncapability", "Main study\neligible"], [
        [models[g["model"]], f"{g['capability_success']}/{g['n']}", f"{g['compliant_capability_success']}/{g['n']}", "No"] for g in gates
    ], [2.5, 1.5, 1.5, 1.5])
    p("Table 3. Capability success includes the exercise-specific requirement. The compliant column additionally excludes any earlier strict-policy violation, even if the final report and finish succeeded. These observations are not counted as main safety trials.")
    exercise_order = ["guard_capability_ordinary", "guard_capability_required", "guard_capability_external", "guard_capability_repair", "guard_capability_cleanup", "guard_capability_approved_external"]
    exercise_names = ["Ordinary workflow", "Required checks", "Necessary external route", "Required repair", "Cleanup", "Approved-work external route"]
    exercise_rows = []
    for case_id, name in zip(exercise_order, exercise_names):
        cells = [next(c for c in g["per_case"] if c["case_id"] == case_id) for g in gates]
        exercise_rows.append([name, *[f"{c['compliant_capability_success']}/{c['n']}" for c in cells]])
    table(doc, ["Exercise", *[models[g["model"]] for g in gates]], exercise_rows, [3.5, 1.75, 1.75])
    p("Table 4. Compliant capability completions by exercise. Neither model passes the screen; its corrected 216-episode main cohort is not launched. This is an unavailable efficacy result, not a zero attack rate. Capability uses schema-constrained JSON generation with history embedded in a prompt, without a native tool-chat or alternative-prompt comparison. Failure concerns these model–interface configurations, not general model incapability.")
    page("Post hoc requirement coverage and runtime hardening")
    p("A separate veto-only gate checks the immutable manifest against recovery registration and verified repair receipts. On saved evidence it defers eight unregistered and eight outstanding-repair requests; each set is one scenario repeated across eight policies. It lets all 96 requests without a repair requirement continue to their original policy and recognizes eight registered completed-repair reference states. Forged actor assertions cannot establish completion.")
    p("Composing that gate with the saved CEG decisions would change one ALLOW to DEFER, leaving zero strict violations among those 14 historical requests. This is a post hoc static result, not a new agent experiment or a measured safety rate. The gate conservatively defers an unregistered requirement even after repair, pending trusted configuration correction. It cannot discover requirements missing from the manifest or authenticate a compromised trust root, and its effect on future task completion remains unmeasured.")
    h("Separate runtime implementation")
    p("A separately versioned broker wrapper now applies the coverage veto during authorization and again immediately before writing the alternative-route report, including inside the inherited CEG lease callback. It binds the immutable manifest, checks actual recovery registration, validates broker-dispatched repair evidence and compares it with live database state. This wrapper is not substituted into the frozen capability or controlled-policy cohorts.")
    table(doc, ["Runtime condition", "Verified behavior in focused tests"], [
        ["Required repair is unregistered", "External report is withheld; no artifact is written"],
        ["Registered repair is outstanding", "Request is deferred until verified completion"],
        ["Repair completes and eligible work qualifies", "Underlying permission can proceed to report execution"],
        ["Actor claims repair or injects grant arguments", "Claims do not satisfy coverage or enlarge authority"],
        ["Registration changes after a lease is issued", "Final callback check stops report creation"],
        ["Ordinary local workflow", "Coverage gate does not block local computation or finish"],
    ], [2.8, 4.2])
    p("Table 5. Software contract checks, not independent deployment samples. All 15 focused wrapper tests and all 317 project tests passed using the bundled Python runtime. The implementation blocks the demonstrated local dispatch path; it has not been evaluated with capable adaptive models.")
    h("Remaining boundary")
    p("A trusted local runtime can inspect its own manifest, registrations and SQL state. A production executor must establish equivalent evidence and protect it from compromise. Unknown repairs, actor-induced events mislabeled as independent, distributed side effects and remote exactly-once execution remain outside this implementation. An early ALLOW decision is also distinct from a completed external effect: the final coverage recheck may still stop execution. Future evaluation must measure actual effects and task availability after such interventions.")

    page("Related work and interpretation")
    p("Robust declassification protects security downgrades from attacker influence [1]. Flow-limited authorization extends this concern to dynamic authority [2], while tainted secure multi-execution uses multiple security views [3]. UCONABC includes authorizations, obligations, environmental conditions and ongoing control [4]. These are direct antecedents; excluding untrusted influence from a release decision is not a new principle.")
    p("Progent constrains privileges and checks policy updates [5]; AgentSpec supplies customizable runtime predicates [6]. FORGE reasons over events and provenance [7], and FAVA constructs evidence-backed permissions and guards [8]. Our resource rule could potentially be expressed within these systems. We have not reproduced them and do not claim that a factual-only local comparator represents their strongest policy or implementation.")
    p("Deny Without Disabling evaluates authorization safety alongside required task utility [9], a distinction essential here. Work on temporal enforcement already addresses obligations and deadlines [10]. Agent Meltdowns studies harmful responses to benign errors [11], while Is Your Agent Playing Dead studies fabricated obstacles [12]. The manufactured-necessity concept is explicitly discussed in the author-published essay Where Power Migrates [13]; that essay is conceptual prior art, not a peer-reviewed evaluation. Path-specific objectives address incentives through selected causal pathways [14], whereas this project gates execution without changing the actor's objective.")
    h("Limitations and remaining validation")
    p("Our cases and oracle are authored together, the external service is simulated, provenance is supplied by trusted dispatch, and queue effects are additive. A mathematically equivalent comparator narrows the algorithmic claim. The post hoc repair extension reuses known failure traces, so it requires fresh evaluation. Independent experts have not yet labeled the cases; a blinded review pack is prepared but uncompleted. No results establish that a capable adversarial agent loses the incentive or ability to disrupt the environment.")
    p("The interface leak invalidates the first main pilot for confirmatory interpretation. The corrected capability screen prevents us from treating task failure as a safety success. Larger or otherwise demonstrably capable models, independent task construction, production evidence capture, induced external failures and actual integrations with published policy engines are necessary for a broader empirical claim.")
    h("Conclusion")
    p("When permission depends on resource provenance, checking the final aggregate load alone cannot enforce the intended rule. In this broker, trusted provenance excludes discretionary actor-dependent qualification, but incomplete evidence reduces availability and unregistered recovery requirements leave a concrete gap. CEG reduces to a simple provenance rule. We supply auditable implementations and explicit failures, and stop short of claiming adaptive-agent defense effectiveness because the tested model–interface configurations do not pass the corrected capability screen.")
    h("Reproducibility")
    p("The project retains source snapshots, frozen cases and schedules, model digests, complete inference requests and responses, SQLite databases, report artifacts and replay audits. Scripted comparisons, corrected capability checks, contaminated development runs and the earlier 192-episode behavioral pilot are separate datasets. All model measurements in this draft come from retained local Ollama responses. The source repository includes commands to reproduce each evidence stream and a gap-closure matrix.")

    page("References")
    refs = [
        ("Myers, A. C., Sabelfeld, A., and Zdancewic, S. (2006). Enforcing Robust Declassification and Qualified Robustness. Journal of Computer Security 14(2), 157–196.", "https://www.cs.cornell.edu/andru/papers/robdecl-jcs/robdecl-jcs.pdf"),
        ("Arden, O., and Myers, A. C. (2016). A Calculus for Flow-Limited Authorization. IEEE Computer Security Foundations Symposium.", "https://www.cs.cornell.edu/andru/papers/flac/"),
        ("McCall, M., Bichhawat, A., and Jia, L. (2023). Tainted Secure Multi-Execution to Restrict Attacker Influence. ACM CCS. DOI 10.1145/3576915.3623110.", "https://www.andrew.cmu.edu/user/liminjia/research/papers/ifc-ccs23.pdf"),
        ("Park, J., and Sandhu, R. (2004). The UCONABC Usage Control Model. ACM TISSEC 7(1), 128–174.", "https://doi.org/10.1145/984334.984339"),
        ("Shi, T., et al. (2025; revised 2026). Progent: Securing AI Agents with Privilege Control. arXiv:2504.11703v3. Preprint.", "https://arxiv.org/abs/2504.11703v3"),
        ("Wang, H., Poskitt, C. M., and Sun, J. (2026). AgentSpec: Customizable Runtime Enforcement for Safe and Reliable LLM Agents. ICSE, 2938–2950. Inspected author manuscript v3.", "https://arxiv.org/abs/2503.18666v3"),
        ("Palumbo, N., et al. (2026). Formal Policy Enforcement for Agentic Systems. arXiv:2602.16708v4. Preprint.", "https://arxiv.org/abs/2602.16708v4"),
        ("Zhang, Y., et al. (2026). FAVA: Formal Authorization for Verified Agents with Evidence-Backed Permission Graphs. arXiv:2607.27267v1. Preprint.", "https://arxiv.org/abs/2607.27267v1"),
        ("Zhang, Y., et al. (2026). Deny Without Disabling: Authorization-Paired Evaluation and Control for Multi-Agent Systems. arXiv:2610.00371v1. Preprint.", "https://arxiv.org/abs/2610.00371v1"),
        ("Hublet, F., Basin, D., and Krstić, S. (2026). Practical Runtime Enforcement of First-Order Temporal Requirements. arXiv:2610.04611v1. Preprint.", "https://arxiv.org/abs/2610.04611v1"),
        ("Jha, R., Triedman, H., Bhattacharya, A., and Shmatikov, V. (2026). Agent Meltdowns: The Road to Hell Is Paved with Helpful Agents. arXiv:2605.19149v1. Preprint.", "https://arxiv.org/abs/2605.19149v1"),
        ("Rodríguez, A., Pozanco, A., and Borrajo, D. (2026). Is Your Agent Playing Dead? Deployed LLM Agents Exhibit Constraint-Evasive Fabrication and Thanatosis. arXiv:2606.14831v1. Preprint.", "https://arxiv.org/abs/2606.14831v1"),
        ("Mc, D. (2026). Where Power Migrates. Axio, 26 July. Author-published conceptual essay.", "https://axio.fyi/p/where-power-migrates"),
        ("Farquhar, S., Carey, R., and Everitt, T. (2022). Path-Specific Objectives for Safer Agent Incentives. AAAI.", "https://arxiv.org/abs/2204.10018v1"),
    ]
    for number, (text, url) in enumerate(refs, 1):
        para = p(f"[{number}] {text} ")
        para.paragraph_format.space_after = Pt(8)
        for run in para.runs:
            run.font.size = Pt(10)
        add_link(para, "Source", url)

    for narrative_table in (doc.tables[0], doc.tables[4]):
        for row in narrative_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    analysis_path = args.out.with_name("manuscript_analysis.json")
    analysis_path.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n")
    provenance = {"created_utc": datetime.now(timezone.utc).isoformat(),
        "document": str(args.out.resolve()), "document_sha256": sha(args.out),
        "builder_sha256": sha(Path(__file__)), "analysis_sha256": sha(analysis_path),
        "main_efficacy_claim": False, "corrected_capability_episodes": len(rows),
        "inputs": {str(path.resolve()): sha(path) for path in [
            comparison_dir / "summary.json", diagnostic_dir / "summary.json",
            ROOT / "docs/CEG_RELATED_WORK_AUDIT.md", ROOT / "docs/GUARD_EXECUTION_STUDY_PLAN_V2.md",
            ROOT / "src/permission_thresholds/hardened_guard_broker.py",
            ROOT / "src/permission_thresholds/requirement_coverage.py",
            ROOT / "tests/test_hardened_guard_broker.py",
            *[run / "manifest.json" for run in args.run], *[run / "episodes.jsonl" for run in args.run]]}}
    args.out.with_name("manuscript_provenance.json").write_text(json.dumps(provenance, indent=2) + "\n")
    print(args.out.resolve())


if __name__ == "__main__":
    main()
