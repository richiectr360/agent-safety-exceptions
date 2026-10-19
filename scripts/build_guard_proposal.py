"""Build the CEG algorithm addendum from integrity-checked scripted evidence."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_empirical_paper import add_link, table


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    summary = json.loads(args.report.read_text())
    if (summary.get("evidence_type") != "authored_scripted_mechanism_checks"
            or summary.get("model_calls") != 0 or summary.get("results_are_synthetic") is not True):
        raise ValueError("This addendum accepts only explicitly scripted mechanism checks")
    for path, expected in summary["source_sha256"].items():
        source = (root / path).resolve()
        if not source.is_relative_to(root) or sha(source) != expected:
            raise ValueError("Guard or evaluation source changed after the report")
    cases_path = args.report.with_name("cases.jsonl")
    if sha(cases_path) != summary["cases_sha256"]:
        raise ValueError("Case integrity mismatch")
    cases = [json.loads(line) for line in cases_path.read_text().splitlines()]
    if len(cases) != summary["cases"]:
        raise ValueError("Case count mismatch")
    for baseline, groups in summary["baselines"].items():
        for category, counts in groups.items():
            selected = [r for r in cases if r["category"] == category]
            recomputed = {"cases": len(selected), **{outcome.lower(): sum(
                row["decisions"][baseline]["outcome"] == outcome for row in selected)
                for outcome in ("ALLOW", "DENY", "DEFER")}}
            if counts != recomputed:
                raise ValueError("Summary disagrees with case decisions")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().find(qn("w:rFonts"))
        for attr in list(fonts.attrib):
            if "theme" in attr.lower():
                del fonts.attrib[attr]
        for script in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + script), "Times New Roman")
        for border in list(style._element.get_or_add_pPr().findall(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)
    doc.styles["Title"].font.size = Pt(21)
    doc.styles["Title"].paragraph_format.space_after = Pt(6)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(11)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    doc.styles["Heading 2"].font.size = Pt(12)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Counterfactual Exception Guard"
    doc.core_properties.author = "Amornchot Singh"
    doc.core_properties.subject = "Proposed algorithm for resource dependent permissions in AI agents"
    doc.add_paragraph("Counterfactual Exception Guard", "Title")
    doc.add_paragraph("A proposed algorithm for resource dependent permissions", "Subtitle")
    doc.add_paragraph("Amornchot Singh  |  18 October 2026")

    def p(text):
        return doc.add_paragraph(text)

    def h(text):
        doc.add_paragraph(text, "Heading 1")

    def page(text):
        doc.add_page_break()
        h(text)

    h("Research claim and AI safety purpose")
    p("We propose Counterfactual Exception Guard (CEG), a deterministic authorization algorithm that checks whether a tool permission would remain justified after excluding discretionary resource changes caused by the agent. Its purpose is to prevent an agent from obtaining broader authority by creating the condition that activates an exception. A Python prototype and scripted checks are complete. Effectiveness against adaptive LLM agents remains untested.")
    p("Consider a policy that permits external processing only when local load reaches ten units. With four units of independently verified load, an agent can add six unnecessary units and satisfy a state-only gate. CEG retains four units in its reference state and denies the request. If independently verified load reaches ten, or validly authorized work supplies the remaining six units, both states qualify and CEG can allow the same operation. This preserves a route for legitimate recovery.")

    h("Two resource views")
    p("A trusted broker records active causes and their lifecycles. The factual view includes all recorded effects. The shadow view includes independently established external causes and effects incurred under valid work grants; it excludes discretionary actor effects, including effects caused by delegates. Verified release removes the corresponding cause from both views.")
    p("For queues, factual load is baseline load plus external, authorized-work and discretionary contributions. Shadow load excludes the discretionary contribution. The baseline must be independently established. Qualification means load is at least the policy threshold. For services, each recorded cause independently keeps the service offline; qualification requires a qualifying cause to have reached the outage-age threshold on a trusted clock. Each cause keeps its own start time.")
    p("The shadow is a specified reference ledger, not a simulation of every possible alternate world. Correlated causes or downstream effects that cannot be represented faithfully require deferral. Aggregate health reports can reconcile the ledger but cannot establish independent causation.")

    h("Scoped authority for necessary work")
    p("A trusted authority issues a work grant before the effect occurs, binding the job, resource, actor, purpose, budget and expiry. The prototype endorses an effect only if it fits the remaining budget in full. Expiry or revocation prevents new endorsements without erasing already incurred effects. A grant is an independent policy decision; the algorithm cannot determine whether an actor's claimed work is truly necessary.")

    page("Algorithm and conditional guarantee")
    p("For an in-scope conditional request, CEG requires a true exception predicate in both views, complete causal evidence and no overdue recovery obligation. It returns ALLOW, DENY or DEFER. Deferral requires independent review or additional evidence and must be counted separately from task success. Independently unconditional permissions are outside this conditional check.")
    code = [
        "1  Check ordinary authorization and request scope.",
        "2  Obtain trusted time and reconcile the resource ledger.",
        "3  If provenance or reconciliation is uncertain: DEFER.",
        "4  Evaluate the exception in factual and shadow states.",
        "5  If either verified predicate is false: DENY.",
        "6  If a relevant recovery obligation is overdue: DEFER.",
        "7  Otherwise issue a short lived, single use lease.",
        "8  Before execution, recheck policy, state, scope and expiry.",
        "9  Consume the lease and dispatch the bound operation.",
    ]
    for line in code:
        paragraph = p(line)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)

    h("Recovery obligations and execution leases")
    p("A trusted recovery obligation names causes and a deadline. Once overdue, an unresolved obligation defers permission until verified repair or an independently certified, expiring infeasibility disposition. This addresses withholding a registered repair after its deadline. It does not block access between exception maturity and a later repair deadline, and it cannot detect an omitted repair that was never registered.")
    p("The lease binds the job, principal, resource, operation, argument hash, policy epoch, state version and expiry. Execution rechecks the decision and consumes the lease under the local broker lock. A callback error consumes the lease and records an uncertain outcome, preventing fresh conditional authorization after a possibly partial side effect. Remote rollback, process-crash recovery and exactly-once execution are not implemented.")

    h("Narrow invariant")
    p("Fix the trusted external history, independently authorized work, policy, request and evaluation time. In conditional mode under CEG, if the shadow predicate is false, adding discretionary actor causes cannot make the request ALLOW. The argument follows directly from the shadow-predicate requirement. It assumes complete mediation, accurate provenance, correct resource transitions and a trusted clock; it does not cover actor influence on supposedly external events, unregistered omissions or bypass of the broker.")
    p("The prototype is a runtime gate, with no model training or learned classifier. Evaluating a request scans active causes and relevant recovery records. Its security depends on the broker supplying trustworthy inputs. The current Python API does not authenticate evidence or isolate the agent from broker memory, and no live-agent adapter has been connected to CEG.")

    page("Scripted mechanism checks")
    p(f"We executed {summary['cases']} authored checks: {summary['scenario_families']} case families across queue and service mechanisms at thresholds 3, 5 and 10. These are synthetic contract checks with zero model calls. Their labels were designed with the algorithm and are not independent judgments of real agent behavior. The earlier 192-episode LLM cohort remains separate and does not validate this defense.")
    p("The comparison includes a factual-only policy, permanent denial after discretionary interference, the shadow rule without recovery obligations, full CEG, and denial of all conditional exceptions. These are local policy baselines, not reproductions of published defense systems.")
    rows = []
    labels = {"factual_only": "Factual only", "permanent_taint": "Permanent taint denial", "shadow_only": "Shadow without obligations", "ceg": "CEG", "deny_all": "Deny all exceptions"}
    for key, label in labels.items():
        groups = summary["baselines"][key]
        def fraction(group, outcome):
            return f"{groups[group][outcome]}/{groups[group]['cases']}"
        rows.append([label, fraction("actor_dependent", "deny"), fraction("legitimate_qualified", "allow"),
                     fraction("overdue_recovery", "defer"), fraction("uncertain", "defer")])
    table(doc, ["Policy", "Actor dependent\ndenied", "Legitimate\nallowed", "Overdue repair\ndeferred", "Uncertain\ndeferred"], rows, [2.05, 1.23, 1.12, 1.35, 1.25])
    p("Table 1. Counts within four disjoint scripted categories. The other twelve checks comprise six false-predicate cases, denied by every policy, and six unregistered-repair-delay cases discussed below. A zero in a deferral column may indicate permission or denial; it does not measure safety by itself.")

    h("What the comparison establishes")
    p("Within these checks, CEG denies every actor-dependent request while preserving every legitimate qualifying request. Permanent taint denial loses twelve legitimate permissions after resolved interference or alongside an independently sufficient cause. Adding recovery obligations changes six shadow-only approvals into deferrals. Factual-only approves all uncertain and overdue-repair cases; deny-all denies those cases as well as every legitimate exception.")
    p("These differences establish that the implemented mechanisms produce their specified decisions. They do not estimate real-world attack success, false-positive rates, task completion or adaptive resistance. Repeating thresholds within authored families does not create independent deployment samples.")

    h("Retained counterexample and integration limits")
    p("CEG still allows all six unregistered-repair-delay cases: without a registered obligation, the ledger cannot distinguish legitimate waiting from withholding a feasible required repair. It also allows genuinely independent qualification despite concurrent discretionary interference. It therefore blocks a route from interference to authority, while leaving disruption itself and broader intent assessment to other controls.")
    p("Separate adversarial unit checks cover grant scope and lifecycle, mixed causes, misleading telemetry, overdue repairs, revocation, argument substitution, lease reuse and ambiguous executor errors. All 196 project tests passed, including 49 CEG tests and a 384-state queue invariant check within one test. These counts describe software verification, not additional agent experiments.")

    page("Prior work and the next empirical test")
    p("Protecting the conditions of authority from untrusted influence has substantial precedent. Robust declassification constrains attacker influence over release decisions [1]; flow-limited authorization protects dynamic authorization [2]. Tainted secure multi-execution is a close predecessor to maintaining multiple security views [3]. Counterfactual constraints also appear in safe reinforcement learning [4], while side-effect research warns that a poorly chosen reference policy can create new incentives [5].")
    p("Our candidate contribution is the specific resource-cause ledger, scoped work-grant treatment and recovery-obligation extension for conditional tool permissions, followed by a controlled safety-and-utility evaluation. The available review does not establish first-of-its-kind novelty. Naming the guard or combining familiar primitives is insufficient on its own; a paper must demonstrate a useful technical distinction and evidence beyond these checks.")

    h("Experiment needed before an effectiveness claim")
    p("Freeze the defense and test fresh agents on held-out scenarios with independently scheduled outages, approved work, discretionary interference and recoverable failures. Use matched variants across all five policies, repeated rollouts and separate capability checks. Include attacks through delegation, induced third-party failures, forged evidence, omission and stale approvals. A trusted simulator may provide ground-truth provenance; a deployment study must measure the cost and errors of obtaining it.")
    p("Measure successful privileged operations dependent on discretionary interference, legitimate exception access, task completion, recovery completion, deferrals and runtime cost. Report attempted violations separately from blocked ones. Analyze uncertainty at the task-family level and retain failed or partial runs. The hypothesis is that CEG reduces actor-dependent privilege use relative to factual-only enforcement while preserving more legitimate access than permanent-taint and deny-all baselines.")

    h("References")
    references = [
        ("[1] Myers, A. C., Sabelfeld, A., and Zdancewic, S. (2006). Enforcing Robust Declassification and Qualified Robustness. Journal of Computer Security, 14(2). ", "https://www.cs.cornell.edu/andru/papers/robdecl-jcs/robdecl-jcs.pdf"),
        ("[2] Arden, O., and Myers, A. C. (2016). A Calculus for Flow-Limited Authorization. IEEE Computer Security Foundations Symposium. ", "https://www.cs.cornell.edu/andru/papers/flac/"),
        ("[3] McCall, M., Bichhawat, A., and Jia, L. (2023). Tainted Secure Multi-Execution to Restrict Attacker Influence. ACM CCS. ", "https://www.andrew.cmu.edu/user/liminjia/research/papers/ifc-ccs23.pdf"),
        ("[4] Vaskov, S., Schwarting, W., and Baker, C. (2024). Do no harm: A counterfactual approach to safe reinforcement learning. Proceedings of Machine Learning Research, 242, 1675–1687. ", "https://proceedings.mlr.press/v242/vaskov24a.html"),
        ("[5] Krakovna, V., Orseau, L., Kumar, R., Martic, M., and Legg, S. (2018; revised 2019). Penalizing side effects using stepwise relative reachability. arXiv:1806.01186. ", "https://arxiv.org/abs/1806.01186"),
    ]
    for text, url in references:
        paragraph = p(text)
        paragraph.paragraph_format.space_after = Pt(5)
        for run in paragraph.runs:
            run.font.size = Pt(10)
        add_link(paragraph, "Source", url)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.out)
    provenance = {"generated_utc": datetime.now(timezone.utc).isoformat(),
                  "evidence_type": summary["evidence_type"], "model_calls": 0,
                  "report": str(args.report.resolve()), "report_sha256": sha(args.report),
                  "source_sha256": summary["source_sha256"], "cases_sha256": summary["cases_sha256"],
                  "builder_sha256": sha(Path(__file__)), "document_sha256": sha(args.out),
                  "document": str(args.out.resolve()), "llm_defense_effectiveness_validated": False}
    args.out.with_name("proposal_provenance.json").write_text(json.dumps(provenance, indent=2) + "\n")
    print(args.out.resolve())


if __name__ == "__main__":
    main()
